# -*- coding: utf-8 -*-
"""Word de plan de trabajo: elección global + jornada en el candidato."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "SIGEVA-Plan-Trabajo-Eleccion-Global.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
GREEN = RGBColor(0x1F, 0x7A, 0x4D)
SENA = RGBColor(0x00, 0x73, 0x32)
MUTED = RGBColor(0x4A, 0x55, 0x63)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0x92, 0x40, 0x0E)

NAVY_HEX = "1B3A4B"
GREEN_HEX = "1F7A4D"
SENA_HEX = "92D050"
ROW_HEX = "F2F4F7"
SOFT_GREEN = "E8F5EE"
SOFT_AMBER = "FEF3C7"
SOFT_NAVY = "E8EEF2"
SOFT_RED = "FDECEC"


def set_run(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(key), "Arial")


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def no_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def set_cell_margins(cell, **sides):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side, cm in sides.items():
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(int(cm * 567)))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def cell_text(cell, text, *, bold=False, color=BLACK, size=10, fill=None, align="left"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)
    set_cell_margins(cell, top=0.08, bottom=0.08, left=0.12, right=0.12)


def add_p(doc, text, *, size=11, bold=False, color=BLACK, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    size = {1: 16, 2: 13, 3: 12}.get(level, 12)
    for run in p.runs:
        set_run(run, size=size, bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_table(doc, headers, rows, col_cm=None, header_fill=NAVY_HEX):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell_text(table.rows[0].cells[i], header, bold=True, color=WHITE, fill=header_fill, size=10)
    for r_i, row in enumerate(rows):
        bg = ROW_HEX if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, fill=bg, size=10)
    if col_cm:
        for row in table.rows:
            for i, width in enumerate(col_cm):
                row.cells[i].width = Cm(width)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)
    return table


def banner(doc, title, subtitle, fill=NAVY_HEX):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=0.28, bottom=0.28, left=0.35, right=0.35)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    set_run(r1, size=14, bold=True, color=WHITE)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(subtitle)
    set_run(r2, size=10, color=WHITE)
    no_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def callout(doc, title, body, fill=SOFT_AMBER, title_color=AMBER):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=0.18, bottom=0.18, left=0.28, right=0.28)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    set_run(r1, size=11, bold=True, color=title_color)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_run(r2, size=10, color=BLACK)
    p2.paragraph_format.space_after = Pt(0)
    no_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def step(doc, number, title, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(f"Paso {number}.  {title}")
    set_run(r, size=11, bold=True, color=GREEN)
    for line in lines:
        add_p(doc, line, size=11, space_after=4)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def header_footer(doc):
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.paragraph_format.space_after = Pt(4)
        r = hp.add_run("SIGEVA")
        set_run(r, size=9, bold=True, color=SENA)
        r2 = hp.add_run("   ·   Plan de trabajo   ·   Elección global")
        set_run(r2, size=9, color=MUTED)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = fp.add_run("Alex  ·  Maicol  ·  Mebel  ·  Paula  ·  Sofia")
        set_run(r, size=8, color=MUTED)
        r2 = fp.add_run("          Pág. ")
        set_run(r2, size=8, color=MUTED)
        add_page_number(fp)


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    header_footer(doc)

    # Portada corta
    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(2)
    r = t.add_run("SIGEVA")
    set_run(r, size=12, bold=True, color=SENA)

    add_p(doc, "Sistema de Gestión Electoral y Validación de Votos", size=11, color=MUTED, space_after=2)

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(4)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run("Plan de trabajo — elección global")
    set_run(r, size=22, bold=True, color=NAVY)

    add_p(
        doc,
        "Instructivo para el equipo. Dónde ir, qué quitar, qué poner y por qué. "
        "Mañana 26 de agosto de 2026. Esto es código, no documentación.",
        size=11,
        color=MUTED,
        space_after=4,
    )
    add_p(
        doc,
        "Alex (back, crear elección)  ·  Maicol (back, candidatos)  ·  Mebel (back, listar y filtrar)  ·  "
        "Paula (front, form elección)  ·  Sofia (front, tarjetón)",
        size=10,
        color=MUTED,
        italic=True,
        space_after=10,
    )

    callout(
        doc,
        "La regla de producto, en una frase",
        "Una elección por centro. La jornada ya no vive en la convocatoria: vive en cada candidato. "
        "Mañana, tarde y noche corren en la misma urna. El aprendiz eligiendo jornada en el primer login no entra mañana.",
        fill=SOFT_GREEN,
        title_color=GREEN,
    )

    # 1. Qué cambia
    add_h(doc, "1. Qué estamos cambiando y por qué")
    add_p(
        doc,
        "Hoy la elección se crea con jornada (Mañana, Tarde o Noche). Eso obliga al funcionario a armar tres "
        "elecciones para el mismo centro, la misma fecha y el mismo cargo. Es el mismo proceso copiado tres veces.",
    )
    add_p(
        doc,
        "La idea nueva: la elección es global al centro (nombre + ventana de fechas y horas). Al inscribir candidatos "
        "se les marca la jornada. El aprendiz, cuando vote, solo ve los candidatos de su jornada. Todo en una sola elección.",
    )
    add_table(
        doc,
        ["Hoy", "Mañana (to-be)"],
        [
            ["3 elecciones (una por jornada)", "1 elección del centro"],
            ["elecciones.jornada es obligatoria", "La elección no lleva jornada"],
            ["El tarjetón es de esa jornada porque la elección lo es", "candidatos.jornada: Mañana, Tarde o Noche"],
            ["El aprendiz lista elecciones filtradas por jornada", "Lista una elección; el GET de candidatos filtra"],
        ],
        col_cm=[8.4, 8.4],
    )

    add_h(doc, "2. Qué pasa con voto, conteo y ganador", 2)
    add_table(
        doc,
        ["Pieza", "¿Se rompe?", "Qué hacemos mañana"],
        [
            [
                "Voto único",
                "No. Sigue siendo un aprendiz, una elección, un voto.",
                "No abrir votoxcandidato ni el middleware.",
            ],
            [
                "Conteo por candidato",
                "No. Cada fila de voto sigue apuntando a un candidato.",
                "No tocar quantityVotes.",
            ],
            [
                "Ganador / acta",
                "Sí, si se deja un solo máximo. Mezclaría las tres jornadas.",
                "No reescribir el PDF. Queda para la siguiente tanda: acta por jornada.",
            ],
        ],
        col_cm=[3.6, 6.6, 6.6],
    )
    callout(
        doc,
        "Cerrar en el daily (2 minutos)",
        "Con este flujo el resultado no puede ser un ganador único de todo el centro. "
        "Mañana, tarde y noche son tres contiendas en la misma elección. El acta se partirá por jornada. "
        "Mañana no se implementa: solo se deja dicho para no pelearlo después.",
    )

    # Orden
    add_h(doc, "3. Orden del día (para no pisarse en git)")
    add_p(
        doc,
        "Alex y Mebel se parten eleccion_controller.ts por método, no por líneas sueltas. "
        "Maicol es dueño de todo lo de candidatos. Paula y Sofia trabajan en el frontend (otro repo).",
    )
    add_table(
        doc,
        ["Cuándo", "Qué", "Quién"],
        [
            ["08:10", "Daily: leer la regla de producto. Nadie improvisa el body del POST.", "Todos"],
            ["Primero", "ALTER TABLE candidatos ADD COLUMN jornada VARCHAR(20);", "Maicol"],
            ["Enseguida", "Quitar jornada de crear/editar elección. DROP NOT NULL en elecciones.jornada.", "Alex"],
            ["En paralelo", "Cambiar los GET: dejar de mandar jornada de la elección; filtrar candidatos.", "Mebel"],
            ["En paralelo", "Form de elección sin select de jornada.", "Paula"],
            ["Cuando el ALTER exista", "Form de candidato + tarjetón con ?jornada=.", "Sofia"],
            ["Tarde", "Humo de 15 minutos (sección 10).", "Los cinco"],
        ],
        col_cm=[3.2, 10.4, 3.2],
    )

    # ALEX
    banner(
        doc,
        "4. Alex  —  backend, crear la elección",
        "Repo sigevaBack   ·   Archivo dueño: eleccion_controller.ts (solo crear y actualizar)",
    )
    add_p(
        doc,
        "Tu trabajo es que se pueda crear una sola convocatoria por centro. Hoy el POST exige jornada y responde 400. "
        "Ese 400 es lo que obliga a tres elecciones.",
    )
    step(
        doc,
        1,
        "app/controllers/eleccion_controller.ts  →  método crearEleccion",
        [
            "Quita 'jornada' de request.only([...]).",
            "Borra el if que compara con Mañana / Tarde / Noche y devuelve 400. No lo comentes: bórralo.",
            "Deja las validaciones de centro, fechas y horas. Esas sí son de la elección.",
            "Eleccione.create(dataEleccion) ya no manda jornada. La columna queda null.",
        ],
    )
    step(
        doc,
        2,
        "El mismo archivo  →  método actualizarEleccion",
        [
            "Lo mismo: saca jornada del only y borra el if idéntico.",
            "No edites jornada de la elección. Ya no es dato de la convocatoria.",
            "No toques desde traerPorCentroFormacion hacia abajo. Eso es de Mebel.",
        ],
    )
    step(
        doc,
        3,
        "app/models/eleccione.ts y SQL",
        [
            "Cambia declare jornada: string  por  string | null. Si no, el create truena.",
            "En PostgreSQL: ALTER TABLE elecciones ALTER COLUMN jornada DROP NOT NULL;",
            "No borres la columna. Las elecciones viejas siguen con valor; las nuevas nacen vacías.",
        ],
    )
    callout(
        doc,
        "Listo cuando",
        "Postman o Thunder: POST /api/eleccion/crear sin jornada, con nombre, centro, fechas y horas, responde 201. "
        "Si mandan jornada de más, se ignora; no revienta.",
        fill=SOFT_GREEN,
        title_color=GREEN,
    )
    callout(
        doc,
        "No abras",
        "candidatos, votos, OTP, FiltroJorElCen, organizacion, centros.",
        fill=SOFT_RED,
        title_color=RGBColor(0x8B, 0x1E, 0x3F),
    )

    # MAICOL
    banner(
        doc,
        "5. Maicol  —  backend, jornada en el candidato",
        "Repo sigevaBack   ·   Eres dueño de modelo, validator, service y controller de candidatos",
    )
    add_p(
        doc,
        "El funcionario mete aprendices de mañana, tarde y noche en la misma elección. Cada fila de candidatos "
        "tiene que saber en cuál jornada corre. Sin este campo Sofia no puede filtrar el tarjetón. "
        "No leas la jornada del grupo: el padrón todavía no trae eso fiable; la pone el funcionario al inscribir.",
    )
    step(
        doc,
        1,
        "SQL primero (bloquea a Sofia y a Mebel)",
        [
            "ALTER TABLE candidatos ADD COLUMN jornada VARCHAR(20);",
            "No pongas NOT NULL el primer minuto si ya hay filas. Los valores válidos son exactamente: Mañana, Tarde, Noche.",
        ],
    )
    step(
        doc,
        2,
        "app/models/candidatos.ts",
        [
            "Agrega @column() declare jornada: string  junto a numero_tarjeton.",
            "No toques el idEleccion duplicado ni Cloudinary.",
        ],
    )
    step(
        doc,
        3,
        "app/validators/store_candidato_validator.ts",
        [
            "En el vine.object agrega jornada como enum de Mañana, Tarde y Noche (con eñe, igual que el crear elección viejo).",
            "Si el front manda “mañana” en minúscula, Vine rechaza. Eso está bien.",
        ],
    )
    step(
        doc,
        4,
        "app/services/candidatos_service.ts",
        [
            "CreateCandidatoDTO: agrega jornada: string.",
            "checkDuplicateTarjeton hoy busca (ideleccion, numero_tarjeton). Eso impide un 01 en mañana y un 01 en tarde. "
            "Cámbialo a filtrar también jornada. Firma: (ideleccion, numero_tarjeton, jornada).",
            "createWithOptionalUpload: pasa jornada al Candidatos.create y llama al unique con los tres argumentos.",
            "getAllCandidatosByIdEleccion: segundo argumento opcional jornada. Si viene, andWhere. Si no, lista los tres bloques. Deja el RANDOM(): eso es la urna.",
            "updateCandidatos: si llega jornada, asígnala. Si cambia tarjetón o jornada, vuelve a correr el unique.",
        ],
    )
    step(
        doc,
        5,
        "app/controllers/candidatos_controller.ts",
        [
            "store: pasa jornada: payload.jornada al service.",
            "update: mete 'jornada' en el request.only.",
            "getByEleccion: lee request.qs().jornada y pásala al service.",
            "Misma ruta. No crees otra: GET /api/candidatos/listar/:ideleccion?jornada=Tarde",
        ],
    )
    callout(
        doc,
        "Listo cuando",
        "Tres POST a /api/candidatos/crear en la misma elección, tarjetón 01, uno por jornada: los tres pasan. "
        "GET sin query = 3. GET ?jornada=Tarde = 1.",
        fill=SOFT_GREEN,
        title_color=GREEN,
    )
    callout(
        doc,
        "No abras",
        "eleccion_controller, votos, OTP, grupo, ImportController.",
        fill=SOFT_RED,
        title_color=RGBColor(0x8B, 0x1E, 0x3F),
    )

    # MEBEL
    banner(
        doc,
        "6. Mebel  —  backend, listar y filtrar (esto es código)",
        "Repo sigevaBack   ·   GET de elección   ·   No es Word, no es HU",
    )
    add_p(
        doc,
        "Alex deja de escribir jornada en la elección. Tú dejas de leerla. Hoy los listados mandan "
        "jornada: eleccion.jornada y FiltroJorElCen filtra por el grupo del aprendiz. Eso es el modelo viejo. "
        "Si Alex crea bien y los GET siguen mandando jornada de la elección, Paula ve tres cards o un select fantasma. "
        "Los listados son la mitad del flujo.",
    )
    step(
        doc,
        1,
        "app/services/FiltroJorElCen.ts  →  filtroElecciones",
        [
            "Quita el preload que entra a aprendiz → grupo y hace where sobre grupo.jornada. Eso filtra la ficha, no el tarjetón.",
            "Deja el filtro de centro y de fechas.",
            "En el preload de candidato: si jornada viene, where sobre candidatos.jornada (la columna de Maicol). Si no viene, trae todos.",
            "La URL no cambia: GET /api/eleccion/traerFiltro  (método traerFiltrado del controller).",
        ],
    )
    step(
        doc,
        2,
        "app/controllers/eleccion_controller.ts  —  solo desde traerPorCentroFormacion hasta el final",
        [
            "traerPorCentroFormacion y traerPorCentroFormacionTodas: en el push del JSON hay jornada: eleccion.jornada. Quítala (o mándala null). No borres fecha, hora, centro, nombre, ideleccion.",
            "Hay un comentario viejo que adivina jornada por el primer candidato. No lo actives.",
            "traerPorJornada (GET /api/eleccionJornada/listar): deja de filtrar grupo. Si llega ?jornada=, filtra candidatos.jornada. No borres la ruta: el front de urna ya la pega. Cambias el significado, no el path.",
        ],
    )
    step(
        doc,
        3,
        "app/services/EleccionesServices.ts y swagger si te alcanza",
        [
            "listarPorCentro ya precarga candidatos. No filtres eleccion.jornada. El JSON del candidato saldrá con jornada cuando Maicol la persista.",
            "start/swagger.ts: en /api/eleccion/crear quita jornada del example. En /api/eleccionJornada/listar el summary pasa a ser “candidatos de una jornada dentro de la elección”, no “elecciones por jornada”.",
        ],
    )
    callout(
        doc,
        "Listo cuando",
        "GET por centro devuelve una elección sin campo jornada (o null). "
        "GET /api/eleccionJornada/listar?jornada=Tarde trae esa elección con candidatos solo de tarde.",
        fill=SOFT_GREEN,
        title_color=GREEN,
    )
    callout(
        doc,
        "No abras",
        "crearEleccion, actualizarEleccion, el PDF del acta, aprendizs_controller, el picker de primer login, ningún .docx de HU.",
        fill=SOFT_RED,
        title_color=RGBColor(0x8B, 0x1E, 0x3F),
    )

    # PAULA
    banner(
        doc,
        "7. Paula  —  frontend, formulario de elección",
        "Repo del front (gestión)   ·   El back de este workspace no trae pantallas",
    )
    add_p(
        doc,
        "Busca la pantalla donde el funcionario crea o edita una elección: el form que tiene el select Mañana / Tarde / Noche "
        "junto a las fechas. Ese select es lo que hay que matar. Si el back ya es global y el form sigue pidiendo jornada, "
        "el funcionario sigue creando tres veces.",
    )
    step(
        doc,
        1,
        "Encuentra el código",
        [
            "Busca en el front: jornada, Mañana, /api/eleccion/crear, /api/eleccionActualizar, eleccionPorCentro.",
        ],
    )
    step(
        doc,
        2,
        "Form crear y editar",
        [
            "Quita el input o select de jornada del JSX y del state.",
            "El POST a /api/eleccion/crear no manda jornada. Body: idcentro_formacion (el del usuario logueado, no un combo de otra sede), nombre, fecha_inicio, fecha_fin, hora_inicio, hora_fin.",
            "Lo mismo en el PUT /api/eleccionActualizar/:idEleccion.",
            "Validación del form: deja nombre y ventana. Quita “jornada obligatoria”.",
        ],
    )
    step(
        doc,
        3,
        "Listado de elecciones del centro",
        [
            "Donde pintas eleccion.jornada en la card, quítalo. Una elección = una fila.",
            "Si el UI agrupa en tres columnas Mañana / Tarde / Noche, eso sale: ahora hay una convocatoria.",
            "Endpoints: /api/eleccionPorCentro/:id, /api/eleccion/traerTodas/:id, /api/eleccion/centrof/:id. Mebel deja de enviar jornada ahí. No filtres en el cliente.",
        ],
    )
    callout(
        doc,
        "Ojo con el merge de Alex",
        "Si pruebas el form antes de su merge, el POST viejo te va a devolver 400 “La jornada no es válida”. "
        "No es un bug tuyo: el back todavía exige el campo.",
        fill=SOFT_NAVY,
        title_color=NAVY,
    )
    callout(
        doc,
        "No abras",
        "El form de candidatos (Sofia). Login de urna. Acta. Backend.",
        fill=SOFT_RED,
        title_color=RGBColor(0x8B, 0x1E, 0x3F),
    )

    # SOFIA
    banner(
        doc,
        "8. Sofia  —  frontend, inscribir candidato y tarjetón",
        "Repo del front (gestión + urna)   ·   La jornada se pone aquí, no en la elección",
    )
    add_p(
        doc,
        "Paula quita jornada de la elección. Tú la pones en el candidato. En gestión se ven las tres jornadas. "
        "En urna, solo la del aprendiz. OTP y voto siguen usando ideleccion: una sola urna, un solo voto.",
    )
    step(
        doc,
        1,
        "Form inscribir candidato (gestión)",
        [
            "Busca el POST a /api/candidatos/crear (multipart: foto, propuesta, tarjetón, idaprendiz, ideleccion).",
            "Agrega un select obligatorio Mañana / Tarde / Noche y mándalo como jornada, con esas grafías. Maicol lo valida con Vine.",
            "El aprendiz se elige del censo del centro (ya existe). No inventes la jornada leyendo el grupo.",
        ],
    )
    step(
        doc,
        2,
        "Lista de candidatos en gestión",
        [
            "GET /api/candidatos/listar/:ideleccion sin query. Agrupa en tres bloques.",
            "Si no hay de una jornada, bloque vacío. No es un error.",
        ],
    )
    step(
        doc,
        3,
        "Tarjetón de la urna (sin pantalla nueva de primer login)",
        [
            "El login del aprendiz ya devuelve data.jornada desde el grupo (aprendizs_controller.login). Si viene un string, llama GET /api/candidatos/listar/:ideleccion?jornada= + ese valor.",
            "Si viene null (el padrón no trajo jornada), no armes “elige tu jornada”. Eso es el último paquete.",
            "Donde hoy listas elecciones filtradas por jornada (/api/eleccion/traerFiltro o /api/eleccionJornada/listar): ahora debe aparecer una elección del centro. El recorte es en candidatos, no en cuántas elecciones hay.",
        ],
    )
    callout(
        doc,
        "Listo cuando",
        "En gestión, tres bloques en la misma elección. En urna, si el login trae Tarde, solo esos candidatos. Sin picker nuevo.",
        fill=SOFT_GREEN,
        title_color=GREEN,
    )
    callout(
        doc,
        "No abras",
        "generar OTP, validarVoto, guardar jornada en aprendiz, el form de crear elección (Paula).",
        fill=SOFT_RED,
        title_color=RGBColor(0x8B, 0x1E, 0x3F),
    )

    # Nadie
    add_h(doc, "9. Archivos que mañana nadie abre")
    add_table(
        doc,
        ["Archivo o módulo", "Por qué se deja"],
        [
            [
                "votoxcandidato_controller.ts",
                "El body del voto no cambia: idcandidato + idaprendiz.",
            ],
            [
                "middleware/validarVoto.ts",
                "Un voto por aprendiz por ideleccion. Con una elección global eso es justo lo que queremos.",
            ],
            [
                "validacionVotoController.ts",
                "OTP no sabe de jornada. total_candidatos queda inflado; no vale el día.",
            ],
            [
                "generacion_reporte_controller.ts",
                "El ganador único mezclaría jornadas. Se parte en otra tanda.",
            ],
            [
                "ImportController.ts / grupo_controller.ts / login aprendiz",
                "Censo y primer login. El picker de jornada del aprendiz es último.",
            ],
            [
                "organizacion.ts, centros, regionales, municipios",
                "Fuera de este flujo.",
            ],
        ],
        col_cm=[6.2, 10.6],
    )

    add_h(doc, "10. Humo de la tarde (15 minutos, los cinco)")
    add_p(doc, "Si esto corre, el día sirvió. Nadie vota. Nadie genera acta.")
    add_table(
        doc,
        ["#", "Qué hacer", "Quién dispara"],
        [
            ["1", "Crear una elección del centro, sin jornada.", "Paula + Alex"],
            ["2", "Inscribir tres aprendices, uno por jornada, tarjetón 01 en cada una.", "Sofia + Maicol"],
            ["3", "GET /api/candidatos/listar/:id sin query → 3 candidatos.", "Maicol"],
            ["4", "GET ...?jornada=Tarde → 1 candidato.", "Sofia"],
            ["5", "GET elecciones del centro → una fila, sin jornada de convocatoria.", "Mebel + Paula"],
        ],
        col_cm=[1.4, 11.4, 4.0],
        header_fill=GREEN_HEX,
    )

    add_h(doc, "11. Fuera de mañana (no se reparte)")
    add_p(
        doc,
        "Primer login del aprendiz eligiendo y guardando jornada. JWT. Reescribir el PDF del acta por jornada. "
        "Migrar elecciones viejas que ya tienen jornada en la convocatoria: conviven; las nuevas nacen globales.",
    )

    add_p(
        doc,
        "Documento generado para el equipo SIGEVA. Uso interno. 25 de agosto de 2026.",
        size=9,
        color=MUTED,
        italic=True,
        space_after=0,
    )

    doc.save(OUT)
    print(f"OK {OUT}")


if __name__ == "__main__":
    build()
