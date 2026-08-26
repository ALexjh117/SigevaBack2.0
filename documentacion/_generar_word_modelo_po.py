# -*- coding: utf-8 -*-
"""Word para el Product Owner: por qué el modelo multi-tenant quedó así."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "SIGEVA-Modelo-Multitenant-para-PO.docx"

NAVY = RGBColor(0x1B, 0x3A, 0x4B)
GREEN = RGBColor(0x2F, 0x6B, 0x4F)
MUTED = RGBColor(0x4A, 0x55, 0x63)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
HEADER_HEX = "1B3A4B"
ROW_HEX = "F2F4F7"


def set_run(run, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:eastAsia"), "Arial")
    rFonts.set(qn("w:cs"), "Arial")


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def cell_text(cell, text, *, bold=False, color=BLACK, size=10, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    if fill:
        shade(cell, fill)


def add_p(doc, text, *, size=11, bold=False, color=BLACK, space_after=8, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run(run, size=16 if level == 1 else 13, bold=True, color=NAVY)
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell_text(table.rows[0].cells[i], h, bold=True, color=RGBColor(255, 255, 255), fill=HEADER_HEX)
    for r_i, row in enumerate(rows):
        bg = ROW_HEX if r_i % 2 == 0 else "FFFFFF"
        for c_i, val in enumerate(row):
            cell_text(table.rows[r_i + 1].cells[c_i], val, fill=bg)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = t.add_run("SIGEVA")
    set_run(r, size=11, bold=True, color=GREEN)
    add_p(doc, "Sistema de Gestión Electoral y Validación de Votos", size=11, color=MUTED, space_after=4)
    h = doc.add_paragraph()
    r = h.add_run("Modelo de datos multi-tenant")
    set_run(r, size=22, bold=True, color=NAVY)
    add_p(
        doc,
        "Documento para el Product Owner. Explica qué se modeló, por qué se agregó cada pieza y por qué este diseño permite que el SENA y otras instituciones usen el mismo producto sin mezclar padrones ni votos.",
        size=11,
        color=MUTED,
        space_after=6,
    )
    add_p(doc, "Audiencia: Product Owner  ·  Tipo: análisis (no es código de producción)  ·  Agosto 2026", size=10, color=MUTED, italic=True)

    add_h(doc, "1. Qué se pidió y qué entregamos")
    add_p(
        doc,
        "Se pidió modelar lo que faltaba para que SIGEVA deje de ser “la app del SENA” y pase a ser una plataforma: centros de formación, votaciones y candidatos, con aislamiento entre organizaciones.",
    )
    add_p(
        doc,
        "Esto no es implementar la urna nueva. Es el mapa de datos para que el producto se pueda aceptar: tablas, dueño de cada fila y reglas que impiden que un colegio vea la elección de otro.",
    )
    add_table(
        doc,
        ["Pedido", "Qué quedó en el modelo", "Qué no se inventó"],
        [
            ["Centros de formación", "Sedes de una organización", "CRUD de regionales como valor de urna"],
            ["Votaciones", "Elección con sede, ventana y dueño", "Segunda vuelta o voto en blanco"],
            ["Candidatos", "Persona del censo + tarjetón único", "Campaña, foros o ranking público"],
        ],
    )

    add_h(doc, "2. El problema de producto (por qué no alcanza lo de hoy)")
    add_p(
        doc,
        "Hoy SIGEVA ya funciona como urna del SENA: el funcionario opera un centro, el aprendiz vota en su sede y jornada, el tarjetón es único por elección y el voto no se debe repetir. Eso es multi-sede. No es multi-tenant.",
    )
    add_p(
        doc,
        "El “dueño” implícito es el centro de formación. No hay una organización. Si mañana entra un colegio, su padrón, sus gestores y sus votos vivirían en las mismas tablas, con el riesgo de listar elecciones ajenas o de que un correo “único en todo el planeta” bloquee a dos instituciones distintas.",
    )
    add_table(
        doc,
        ["Hoy (as-is)", "Consecuencia de producto"],
        [
            ["No existe la tabla organización", "No hay cliente dueño de sus datos. Solo hay sedes SENA."],
            ["El centro es el mundo", "Un colegio no encaja: no tiene “centros SENA” ni regionales."],
            ["La tabla se llama aprendiz", "Un colegio no tiene aprendices; sí tiene un censo de votantes."],
            ["El correo es único global", "Dos instituciones no pueden registrar el mismo correo."],
            ["El voto único vive en el middleware", "La urna no está defendida en la base si alguien se salta la API."],
        ],
    )

    add_h(doc, "3. Qué significa multi-tenant aquí")
    add_p(
        doc,
        "Multi-tenant = el mismo software, la misma base, muchas organizaciones. Cada una ve solo su censo, sus sedes, sus elecciones, sus candidatos y sus votos.",
    )
    add_p(
        doc,
        "No significa una base por cliente. No significa una tabla “aprendiz”, otra “estudiante” y otra “empleado”. Eso duplica el producto y rompe la urna: el motor electoral (censo → convocatoria → tarjetón → voto único → acta) es el mismo para todos.",
        space_after=8,
    )
    add_p(doc, "La decisión de diseño", size=12, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Una base, un esquema, una columna de dueño (organización) en todo lo electoral. El SENA es el primer cliente: una organización con muchos centros. Un colegio es otra organización, con una o varias sedes. El código de la urna no se copia.",
    )
    add_p(
        doc,
        "Esa columna no la elige el usuario en un formulario. Sale de la sesión del gestor o del censo del votante. Si alguien manda el id de otra institución, el sistema lo ignora.",
    )

    add_h(doc, "4. Por qué es un buen modelado")
    add_table(
        doc,
        ["Criterio", "Cómo se cumple"],
        [
            ["Un dueño claro", "Toda fila electoral sabe a qué organización pertenece."],
            ["El SENA no se tira", "Regional, ficha y jornada quedan como extras del primer tenant, no como el núcleo."],
            ["Otros clientes caben", "Sede + votante + elección + candidato + voto sirven a un colegio sin tablas nuevas."],
            ["La urna no se diluye", "Un votante, una elección, un voto; OTP no es el voto; la ventana sigue siendo oficial."],
            ["Escala sin copiar código", "Alta de una organización nueva, no un deploy nuevo por cliente."],
            ["Se puede defender", "Unicidades en base (tarjetón, un voto) además de las reglas de negocio."],
        ],
    )
    add_p(
        doc,
        "Si el modelo fuera “una app por institución”, SIGEVA volvería a ser un desarrollo a medida. Si el modelo fuera “el centro es el tenant”, un colegio de una sola sede y el SENA nacional no caben en la misma idea de producto.",
        italic=True,
        color=MUTED,
    )

    add_h(doc, "5. Cómo se traduce el lenguaje SENA")
    add_p(
        doc,
        "En pantalla, el SENA puede seguir diciendo aprendiz, funcionario y centro. En el modelo de plataforma esas cosas tienen nombre genérico, para no soldar el producto a un solo cliente.",
    )
    add_table(
        doc,
        ["En el SENA (hoy)", "En la plataforma (to-be)", "Por qué se traduce"],
        [
            ["El mundo es el SENA", "Organización (tenant)", "Sin esta fila no hay segundo cliente."],
            ["Centro de formación", "Sede / circunscripción", "Un colegio tiene sedes, no la red SENA."],
            ["Aprendiz", "Votante del censo", "Todos necesitan padrón; no todos forman aprendices."],
            ["Funcionario", "Gestor de una sede", "Es el jurado digital de una urna, no de la red."],
            ["Administrador único", "Admin del tenant + operador de plataforma", "Quien gobierna un cliente no es quien da de alta clientes."],
            ["Jornada Mañana / Tarde / Noche", "Filtro opcional de circunscripción", "Obligatoria en SENA; un colegio puede no usarla."],
        ],
    )

    add_h(doc, "6. Cada pieza: qué es, por qué existe, por qué así")

    add_p(doc, "6.1 Organización (tabla nueva)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Es el tenant: el cliente dueño de los datos. SENA, un colegio, una universidad. Sin esta tabla el sistema no sabe de quién es un padrón. Se agregó porque era el hueco: hoy el centro hace de mundo y eso no escala a otro tipo de institución.",
    )
    add_p(
        doc,
        "Por qué así: un registro por cliente, con nombre, identificador estable (slug) y estado. Si está suspendida, nadie de esa organización entra a urna ni a gestión. El primer registro es el SENA; los centros que ya existen se cuelgan de ese id.",
    )

    add_p(doc, "6.2 Centro de formación / sede (tabla que ya existía, con dueño)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Sigue siendo la circunscripción: ahí vive el censo, el gestor y la urna. No se eliminó porque el SENA la necesita y un colegio también tiene al menos una sede. Lo que cambió es el significado: deja de ser el tenant y pasa a ser una sede de la organización.",
    )
    add_p(
        doc,
        "Por qué así: lleva el id de la organización. La regional SENA queda opcional (un colegio no la usa). El código de sede es único dentro del cliente, no en todo el planeta: dos instituciones pueden tener una sede “01”. El funcionario no elige otra sede; la sesión trae la suya.",
    )

    add_p(doc, "6.3 Votante — censo (hoy se llama aprendiz en el código)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Es el padrón. Toda elección institucional necesita personas identificables: documento, correo, sede, estado. En el SENA esa persona es el aprendiz. En un colegio, el estudiante. En una empresa, el colaborador.",
    )
    add_p(
        doc,
        "Por qué no se borra ni se parte en tres tablas: el motor de la urna es el mismo. Si hubiera aprendiz + estudiante + empleado, cada historia de voto se triplicaría y el acta se volvería inmantenible. Una sola tabla de censo, con id de organización. Ficha, programa y jornada son campos opcionales del SENA; en otro cliente van vacíos.",
    )
    add_p(
        doc,
        "El correo y el documento pasan a ser únicos por organización, no globales. Dos clientes pueden tener un “ana@correo.com” sin chocar.",
    )

    add_p(doc, "6.4 Grupo / ficha (se conserva, no es el núcleo)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "En el SENA la ficha y la jornada (Mañana, Tarde, Noche) definen quién puede entrar a qué urna. Un colegio puede no tener fichas. Por eso el grupo queda como dimensión SENA, no como requisito de la plataforma. El votante puede no traer grupo; el SENA sí lo usa.",
    )

    add_p(doc, "6.5 Usuarios / gestor (tabla que ya existía, con dueño)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Es quien arma el padrón, abre la elección, inscribe candidatos y baja el acta. Se conservó porque esa operación ya existe. Se le agrega organización y se mantiene el centro: el gestor de una sede no opera la red ni la urna de al lado. El administrador del tenant ve las sedes de su institución, no las de otra.",
    )

    add_p(doc, "6.6 Elecciones / votaciones (tabla que ya existía, con dueño)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Es la convocatoria oficial: nombre, sede, ventana de fecha y hora, jornada cuando aplica. Se conservó porque ya es el corazón de SIGEVA. Se le pone organización copiada de la sede, para no filtrar “todas las activas del país”.",
    )
    add_p(
        doc,
        "Por qué no hay un campo “activa”: la urna existe solo si ahora está dentro de la ventana. Guardar un sí/no se desincroniza con el reloj. El aprendiz (votante) solo lista las de su sede, su jornada y ese momento. El gestor lista el inventario de su sede, incluidas las cerradas, para el acta.",
    )

    add_p(doc, "6.7 Candidatos (tabla que ya existía, con dueño)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "No es un nombre suelto. Es una persona del censo de esa organización, inscrita en una elección, con foto, propuesta y número de tarjetón. Se conservó porque el tarjetón ya es el producto. Se exige que el votante sea de la misma sede que la elección.",
    )
    add_p(
        doc,
        "Por qué las unicidades: no puede haber dos “03” en la misma urna, y una persona no se inscribe dos veces en la misma contienda. Eso ya lo pide el producto; el modelo lo deja escrito en la base.",
    )

    add_p(doc, "6.8 Validación de voto / OTP (se limpia, no se inventa otro factor)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Prueba que el votante tiene el correo del censo. Validar el código no registra el voto: son dos actos. Hoy el estado va mezclado en el código (prefijos USED_ / VOTED_). El modelo to-be separa el secreto (6 caracteres) del estado (generado, usado, votado) y le pone organización. No se agregó biometría ni un segundo producto de identidad.",
    )

    add_p(doc, "6.9 Voto (tabla que ya existía, reforzada)", size=13, bold=True, color=NAVY, space_after=4)
    add_p(
        doc,
        "Un sufragio: una persona, una elección, un candidato. Hoy la unicidad vive sobre todo en el middleware. El modelo pide además la unicidad en base (persona + elección) y el id de organización y de elección en la misma fila, para que el acta se reconstruya sin ambigüedad y nadie vote dos veces aunque se salte una capa.",
    )

    add_h(doc, "7. Cómo se aísla de verdad (regla para aceptar el modelo)")
    add_p(doc, "Toda tabla electoral lleva la organización. Al guardar una fila tienen que cumplirse, en cadena:")
    add_table(
        doc,
        ["#", "Regla"],
        [
            ["1", "La sede pertenece a la organización de la sesión."],
            ["2", "El votante pertenece a esa sede y a esa organización."],
            ["3", "La elección pertenece a esa sede y a esa organización."],
            ["4", "El candidato es un votante de la misma sede que la elección."],
            ["5", "El voto apunta a un candidato de esa elección."],
            ["6", "Ese votante no tiene ya un voto en esa elección."],
        ],
    )
    add_p(
        doc,
        "Si alguna igualdad se rompe, no se persiste. Eso es el tenant. No es un filtro de reporte: es la geometría del derecho al voto (sede, y en el SENA también jornada) más el dueño institucional.",
    )

    add_h(doc, "8. Qué queda igual a propósito")
    add_p(
        doc,
        "No se rediseñó la urna. Se le puso dueño. Siguen valiendo las cinco invariantes de SIGEVA: un votante un voto; el OTP caduca y validar no es votar; nadie vota en otra sede ni otra jornada; fuera de la ventana no hay urna; el acta se arma con los votos guardados, incluyendo ceros y empate.",
    )
    add_p(
        doc,
        "Geografía Colombia (departamento, municipio) puede seguir compartida: no es secreto electoral. Regional es catálogo SENA, no el núcleo.",
    )

    add_h(doc, "9. Qué le pedimos aceptar")
    add_table(
        doc,
        ["Afirmación", "Qué implica"],
        [
            ["El tenant es la organización, no el centro", "El SENA es el primer cliente, con muchos centros."],
            ["El censo es una sola tabla de votantes", "No habrá tabla estudiante/empleado. El SENA sigue viendo “aprendiz” en pantalla."],
            ["Ficha y jornada son del SENA", "Otro cliente no está obligado a usarlas."],
            ["El candidato sale del censo", "No se inscribe un nombre que no esté en el padrón de esa sede."],
            ["Este documento es análisis", "No se migró la base todavía. El código actual sigue siendo el as-is SENA."],
        ],
    )
    add_p(
        doc,
        "Cuando este mapa esté aceptado, el orden de construcción (otro ciclo, no este documento) es: crear organización y sembrar SENA; colgar los centros; colgar censo y gestores; colgar elecciones y candidatos; reforzar el voto único en base; separar el estado del OTP. Hasta entonces el producto operativo sigue siendo la red SENA multi-sede.",
    )

    add_h(doc, "10. Dónde está el dibujo")
    add_p(
        doc,
        "El diagrama para Lucid es el archivo Draw.io SIGEVA-modelo-multitenant.drawio, en esta misma carpeta de documentación. Importar: Lucid → Importar documentos → Buscar → ese archivo.",
    )
    add_p(
        doc,
        "Este Word es la explicación. El Draw.io es la foto. Juntos son el entregable de modelado para el Product Owner.",
        italic=True,
        color=MUTED,
    )

    doc.save(OUT)
    print(f"OK {OUT}")


if __name__ == "__main__":
    build()
