# -*- coding: utf-8 -*-
"""Genera entregables/sprint-1 y borra los documentos viejos del análisis."""
from __future__ import annotations

import re
import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(HERE))

from _excel_senior_data_admin import HU_ADMIN  # noqa: E402
from _excel_senior_data_epicas import EPICAS  # noqa: E402
from _excel_senior_data_aprendiz import HU_APRENDIZ  # noqa: E402
from _excel_senior_data_funcionario import HU_FUNCIONARIO  # noqa: E402

OUT = REPO / "entregables" / "sprint-1"
SRC_WORD = HERE / "_plantilla-sena-casos-de-uso.docx"

SENA = "92D050"
NAVY = "1B3A4B"
WHITE = "FFFFFF"
BLACK = "000000"
GOLD = "9A6B2F"
GRAY = "F2F4F7"
LIGHT = "E8F5EE"
LINE = "D0D5DD"
THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
W_NS = "{http://www.w3.org/XML/1998/namespace}"

RNF = [
    ["RNF-NU-001", "Debe tener", "Secretos", "Las claves se guardan con bcrypt. El código de validación tiene 6 caracteres. En producción el código no viaja en la respuesta JSON."],
    ["RNF-NU-002", "Debe tener", "Ventana", "La urna solo existe dentro de la fecha y la hora oficiales de la elección."],
    ["RNF-NU-003", "Debería tener", "Camino corto", "El recorrido del aprendiz es: elecciones, tarjetón, código, voto y comprobante."],
    ["RNF-NU-004", "Debe tener", "Menús", "Cada rol ve solo su menú. El aprendiz no administra. El funcionario no ve la red."],
    ["RNF-NU-005", "Debe tener", "Correo", "Si falla el envío del correo, no queda un voto registrado."],
]


def afont(size=10, bold=False, color=NAVY):
    return Font(name="Arial", size=size, bold=bold, color=color)


def fill(color):
    return PatternFill("solid", fgColor=color)


def align(v="top", h="left"):
    return Alignment(wrap_text=True, vertical=v, horizontal=h)


def getv(data, *keys, default=""):
    for key in keys:
        if isinstance(data, dict) and key in data and data[key] not in (None, ""):
            return data[key]
    return default


def get_list(data, *keys):
    value = getv(data, *keys, default=[])
    if isinstance(value, list):
        return value
    if value:
        return [value]
    return []


def bullets(value):
    if isinstance(value, (list, tuple)):
        return "\n".join(f"• {item}" for item in value)
    return str(value or "")


def prioridad(value):
    text = (value or "").upper()
    if "WON" in text or "NO TENDR" in text:
        return "No tendrá"
    if "SHOULD" in text or "DEBERÍA" in text or "DEBERIA" in text:
        return "Debería tener"
    if "COULD" in text or "PODRÍA" in text or "PODRIA" in text:
        return "Podría tener"
    return "Debe tener"


def prio_colors(value):
    if value == "Debe tener":
        return SENA, BLACK
    if value == "Debería tener":
        return GOLD, WHITE
    if value == "Podría tener":
        return "667085", WHITE
    return "8B1E3F", WHITE


def prio_word(value):
    return "Alta" if prioridad(value) == "Debe tener" else "Media"


def as_ca(item):
    if not isinstance(item, dict):
        return {"id": "", "tipo": "", "dado": str(item), "cuando": "", "entonces": ""}
    return {
        "id": getv(item, "id", "id", "cid"),
        "tipo": getv(item, "tipo", "tipo"),
        "dado": getv(item, "dado", "dado"),
        "cuando": getv(item, "cuando", "cuando"),
        "entonces": getv(item, "entonces", "entonces"),
    }


def ca_cerrados(items):
    cerrados = []
    for raw in items or []:
        item = as_ca(raw)
        tipo = (item.get("tipo") or "").lower()
        if any(flag in tipo for flag in ("pendiente", "condicion", "bloqueado", "no implementar")):
            continue
        cerrados.append(item)
    return cerrados


def criterios_texto(hu):
    lines = []
    for item in ca_cerrados(hu.get("ca")):
        entonces = (item.get("entonces") or "").strip().rstrip(".")
        if entonces:
            lines.append("• " + entonces[0].upper() + entonces[1:] + ".")
    return "\n".join(lines)


def gherkin_corto(hu):
    bloques = []
    for item in ca_cerrados(hu.get("ca")):
        bloques.append(
            f"{item.get('id', '')}\n"
            f"Dado {item.get('dado', '')}\n"
            f"Cuando {item.get('cuando', '')}\n"
            f"Entonces {item.get('entonces', '')}"
        )
    return "\n\n".join(bloques)


def limpia_campo(texto):
    value = texto or ""
    for sep in ("NUNCA", "PROHIBIDO", "⚠️", "PENDIENTE", "PROPUESTO"):
        if sep in value:
            value = value.split(sep)[0]
    return value.strip(" ·.")


def entradas_salidas(hu):
    campos = hu.get("campos") or ""
    partes = re.split(r"salida\s*:", campos, flags=re.I)
    entradas = limpia_campo(partes[0] if partes else campos)
    salidas = limpia_campo(partes[1] if len(partes) > 1 else hu.get("post") or "")
    return entradas or "Datos de la sesión y de la acción.", salidas or (hu.get("post") or "Confirmación de la operación.")


def norm(raw):
    hu = {
        "id": getv(raw, "id", "id"),
        "titulo": getv(raw, "titulo", "titulo"),
        "como": getv(raw, "como", "como"),
        "quiero": getv(raw, "quiero"),
        "para": getv(raw, "para", "para"),
        "actor": getv(raw, "actor"),
        "producto": getv(raw, "producto"),
        "epica": getv(raw, "epica", "epica"),
        "epica_nom": getv(raw, "epica_nom", "epica_nom"),
        "rf": getv(raw, "rf", "rf"),
        "cu": getv(raw, "cu", "cu"),
        "prioridad": prioridad(getv(raw, "moscow", "moscow")),
        "moscow_raw": getv(raw, "moscow", "moscow"),
        "sp": getv(raw, "sp", "sp", default=0),
        "rank": getv(raw, "rank", "rank", default=99),
        "fase": getv(raw, "fase", "fase"),
        "mvp": getv(raw, "mvp", "mvp"),
        "in_scope": get_list(raw, "in_scope", "in_scope", "in_scope"),
        "out_scope": get_list(raw, "out_scope", "out_scope", "out_scope"),
        "pre": getv(raw, "pre", "pre"),
        "post": getv(raw, "post", "post"),
        "campos": getv(raw, "campos", "campos"),
        "deps": getv(raw, "deps", "deps"),
        "bloquea": getv(raw, "bloquea", "bloquea"),
        "dueno": getv(raw, "dueno", "dueno"),
        "ca": get_list(raw, "ca", "ca"),
    }
    hu["historia"] = f"Como {hu['como']}, quiero {hu['quiero']}, para {hu['para']}."
    hu["desc_rf"] = f"El sistema deberá permitir {hu['como']} {hu['quiero']}, para {hu['para']}."
    try:
        hu["sp"] = int(hu["sp"] or 0)
    except Exception:
        hu["sp"] = 0
    try:
        hu["rank"] = int(hu["rank"] or 99)
    except Exception:
        hu["rank"] = 99
    return hu


def all_hus():
    hus = [norm(item) for item in HU_APRENDIZ + HU_FUNCIONARIO + HU_ADMIN]
    ids = [hu["id"] for hu in hus]
    if len(hus) != 22:
        raise SystemExit(f"Se esperaban 22 historias, hay {len(hus)}: {ids}")
    if len(ids) != len(set(ids)):
        raise SystemExit("Hay identificadores duplicados")
    return hus


def paint(cell, value, bg=None, font=None, h="left"):
    if value is not None:
        cell.value = value
    if bg:
        cell.fill = fill(bg)
    cell.font = font or afont()
    cell.alignment = align("top", h)
    cell.border = THIN
    return cell


def widths(ws, sizes):
    for index, size in enumerate(sizes, 1):
        ws.column_dimensions[get_column_letter(index)].width = size


def header_row(ws, headers, bg=SENA, font_color=BLACK):
    for index, title in enumerate(headers, 1):
        paint(ws.cell(1, index), title, bg, afont(10, True, font_color), "center")
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = "SIGEVA  ·  SENA"
    ws.oddHeader.right.text = "&A"


def add_rows(ws, rows, height=48):
    for row_idx, values in enumerate(rows, 2):
        bg = LIGHT if row_idx % 2 == 0 else WHITE
        for col_idx, value in enumerate(values, 1):
            paint(ws.cell(row_idx, col_idx), value, bg)
        ws.row_dimensions[row_idx].height = height
    last = 1 + max(len(rows), 1)
    cols = ws.max_column
    ws.auto_filter.ref = f"A1:{get_column_letter(cols)}{last}"
    try:
        name = re.sub(r"[^A-Za-z0-9]", "", ws.title)[:20] + "Tbl"
        table = Table(displayName=name, ref=f"A1:{get_column_letter(cols)}{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    except Exception:
        pass
    return last


def color_prio(ws, col, start, end):
    for row in range(start, end + 1):
        bg, fg = prio_colors(ws.cell(row, col).value)
        ws.cell(row, col).fill = fill(bg)
        ws.cell(row, col).font = afont(9, True, fg)
        ws.cell(row, col).alignment = align("center", "center")


def kv_sheet(ws, pairs, title):
    ws.sheet_properties.tabColor = SENA
    widths(ws, [36, 92])
    paint(ws.cell(1, 1), title, SENA, afont(16, True, BLACK))
    paint(ws.cell(1, 2), "", SENA)
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 32
    row = 3
    for label, value in pairs:
        paint(ws.cell(row, 1), label, LIGHT, afont(10, True, NAVY))
        paint(ws.cell(row, 2), value, WHITE, afont(10))
        ws.row_dimensions[row].height = 36
        row += 1
    ws.sheet_view.showGridLines = False


def build_sprint():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint 1"
    kv_sheet(
        ws,
        [
            ("Producto", "SIGEVA — Sistema de Gestión Electoral y Validación de Votos"),
            ("Sprint", "Sprint 1 · 24 al 28 de agosto de 2026"),
            ("Qué se hace", "Especificar el producto. En este sprint no se construye código nuevo."),
            ("Qué se entrega", "Requerimientos, casos de uso, historias de usuario y backlog."),
            ("Quién acepta", "Henry y Jorge"),
            ("Equipo", "Alex, Sofia, Paula, Maicol y Mebel"),
            ("Formato de casos de uso", "Plantilla SENA: letra Arial y color verde 92d050."),
            ("Idioma", "Todo el paquete va en español. Prioridad: Debe tener / Debería tener."),
        ],
        "Sprint 1 — qué entregamos",
    )

    ws2 = wb.create_sheet("Entregables")
    ws2.sheet_properties.tabColor = NAVY
    header_row(ws2, ["Orden", "Entregable", "Archivo", "Qué contiene"], NAVY, WHITE)
    add_rows(
        ws2,
        [
            ["1", "Requerimientos", "02-Requerimientos.xlsx", "Requisitos funcionales (uno por historia) y requisitos no funcionales."],
            ["2", "Casos de uso", "03-Casos-de-Uso.docx", "Documento SENA: una tabla por caso de uso, letra Arial, verde 92d050."],
            ["3", "Historias de usuario", "04-Historias-de-Usuario.xlsx", "Un cuadro. Una fila = una historia (Como / quiero / para, qué entra y qué no, criterios)."],
            ["4", "Backlog", "05-Backlog.xlsx", "Las 22 historias de producto, en el orden de entrega de un centro piloto."],
        ],
        40,
    )
    widths(ws2, [10, 24, 36, 88])
    return wb


def build_requerimientos(hus):
    wb = Workbook()
    ws = wb.active
    ws.title = "Requisitos funcionales"
    ws.sheet_properties.tabColor = SENA
    header_row(ws, ["RF", "CU", "Historia", "Nombre", "El sistema deberá", "Actor", "Prioridad", "Épica"])
    rows = []
    for hu in sorted(hus, key=lambda item: item["id"]):
        rows.append([
            hu["rf"], hu["cu"], hu["id"], hu["titulo"], hu["desc_rf"],
            hu["actor"], hu["prioridad"], f"{hu['epica']} {hu['epica_nom']}".strip(),
        ])
    last = add_rows(ws, rows, 56)
    color_prio(ws, 7, 2, last)
    widths(ws, [12, 10, 18, 36, 55, 22, 16, 28])

    ws2 = wb.create_sheet("Requisitos no funcionales")
    ws2.sheet_properties.tabColor = GOLD
    header_row(ws2, ["RNF", "Prioridad", "Tema", "El sistema deberá"], GOLD, WHITE)
    last2 = add_rows(ws2, RNF, 40)
    color_prio(ws2, 2, 2, last2)
    widths(ws2, [14, 16, 16, 90])

    ws3 = wb.create_sheet("Como leer")
    kv_sheet(
        ws3,
        [
            ("Regla", "Una historia de producto = un requisito funcional = un caso de uso."),
            ("Debe tener", "Sin esto no hay elección de un centro."),
            ("Debería tener", "Mejora la operación; el piloto de un centro puede vivir si se pospone."),
            ("Casos de uso", "El detalle SENA (entradas, salidas, criterios y rol) está en 03-Casos-de-Uso.docx."),
            ("No hay CU-08, CU-09 ni CU-19", "La numeración sigue las tarjetas del producto (01-07, 10-18, 20-25)."),
        ],
        "Requerimientos SIGEVA",
    )
    return wb


def build_historias(hus):
    wb = Workbook()
    ws_l = wb.active
    ws_l.title = "Como leer"
    kv_sheet(
        ws_l,
        [
            ("Hoja Historias", "Ahí está el producto. Una fila por historia. Use el filtro."),
            ("Hoja Epicas", "Las capacidades del producto y el paquete de documentos de este sprint."),
            ("Prioridad", "Debe tener = no hay elección de un centro sin eso. Debería tener = puede esperar."),
            ("OTP y voto", "Pedir el código, validarlo y votar son tres historias distintas."),
            ("Dos importaciones", "El funcionario carga el censo de SU centro. El administrador elige el centro. No se fusionan."),
        ],
        "Historias de usuario SIGEVA",
    )

    ws_e = wb.create_sheet("Epicas")
    ws_e.sheet_properties.tabColor = NAVY
    header_row(
        ws_e,
        ["ID", "Épica", "Tipo", "Como / quiero / para", "Problema", "Valor", "Éxito", "Sí entra", "No entra", "Historias", "Prioridad"],
        NAVY,
        WHITE,
    )
    e_rows = []
    for epica in EPICAS:
        e_rows.append([
            getv(epica, "id", "id"),
            getv(epica, "nombre", "nombre"),
            getv(epica, "tipo", "tipo"),
            f"Como {getv(epica, 'como', 'como')}, quiero {getv(epica, 'quiero')}, para {getv(epica, 'para', 'para')}.",
            getv(epica, "problema", "problema"),
            getv(epica, "valor", "valor"),
            getv(epica, "exito", "exito", "exito"),
            getv(epica, "alcance", "alcance"),
            getv(epica, "fuera", "fuera"),
            getv(epica, "historias", "historias"),
            prioridad(getv(epica, "moscow", "moscow")),
        ])
    last_e = add_rows(ws_e, e_rows, 70)
    color_prio(ws_e, 11, 2, last_e)
    widths(ws_e, [14, 26, 28, 42, 36, 32, 36, 40, 40, 28, 16])

    ws = wb.create_sheet("Historias")
    ws.sheet_properties.tabColor = SENA
    header_row(
        ws,
        [
            "ID", "Título", "Historia", "Actor", "Producto", "Épica",
            "RF", "CU", "Prioridad", "Puntos", "Sprint",
            "Sí entra", "No entra", "Criterios de aceptación", "Depende de", "Responsable",
        ],
    )
    rows = []
    for hu in sorted(hus, key=lambda item: item["id"]):
        rows.append([
            hu["id"], hu["titulo"], hu["historia"], hu["actor"], hu["producto"],
            f"{hu['epica']} {hu['epica_nom']}".strip(),
            hu["rf"], hu["cu"], hu["prioridad"], hu["sp"],
            "Sprint 1 — especificar",
            bullets(hu["in_scope"]), bullets(hu["out_scope"]),
            gherkin_corto(hu), hu["deps"], hu["dueno"],
        ])
    last = add_rows(ws, rows, 92)
    color_prio(ws, 9, 2, last)
    widths(ws, [16, 34, 42, 18, 12, 26, 12, 10, 14, 8, 22, 36, 36, 48, 22, 18])
    ws.freeze_panes = "C2"
    return wb


def build_backlog(hus):
    wb = Workbook()
    ws_l = wb.active
    ws_l.title = "Como leer"
    kv_sheet(
        ws_l,
        [
            ("Orden", "Es el orden para armar un centro piloto, no el orden de IDs."),
            ("Este sprint", "Sprint 1 entrega la especificación de estas 22 historias, no el código."),
            ("Debe tener", "Sin eso no hay urna de un centro."),
            ("Archivo de sprint", "01-Sprint.xlsx dice qué se entrega esta semana."),
        ],
        "Backlog de producto SIGEVA",
    )

    ws = wb.create_sheet("Backlog")
    ws.sheet_properties.tabColor = SENA
    header_row(
        ws,
        [
            "Orden", "ID", "Título", "Historia", "Actor", "Producto", "Épica",
            "Prioridad", "Puntos", "Corte", "MVP un centro", "Depende de", "RF", "CU", "Responsable",
        ],
    )
    rows = []
    for hu in sorted(hus, key=lambda item: (item["rank"], item["id"])):
        rows.append([
            hu["rank"], hu["id"], hu["titulo"], hu["historia"], hu["actor"], hu["producto"],
            f"{hu['epica']} {hu['epica_nom']}".strip(),
            hu["prioridad"], hu["sp"], hu["fase"], hu["mvp"], hu["deps"], hu["rf"], hu["cu"], hu["dueno"],
        ])
    last = add_rows(ws, rows, 52)
    color_prio(ws, 8, 2, last)
    widths(ws, [10, 16, 34, 42, 18, 12, 26, 14, 8, 28, 14, 22, 12, 10, 18])
    ws.freeze_panes = "C2"
    return wb


def uniq(row):
    seen = set()
    cells = []
    for cell in row.cells:
        marker = id(cell._tc)
        if marker in seen:
            continue
        seen.add(marker)
        cells.append(cell)
    return cells


def set_cell_text(cell, text):
    tc = cell._tc
    first_p = tc.find(qn("w:p"))
    p_pr = deepcopy(first_p.find(qn("w:pPr"))) if first_p is not None and first_p.find(qn("w:pPr")) is not None else None
    first_r = tc.find(".//" + qn("w:r"))
    r_pr = deepcopy(first_r.find(qn("w:rPr"))) if first_r is not None and first_r.find(qn("w:rPr")) is not None else None
    for paragraph in list(tc.findall(qn("w:p"))):
        tc.remove(paragraph)
    for line in str(text or "").split("\n") or [""]:
        paragraph = OxmlElement("w:p")
        if p_pr is not None:
            paragraph.append(deepcopy(p_pr))
        run = OxmlElement("w:r")
        if r_pr is not None:
            run.append(deepcopy(r_pr))
        node = OxmlElement("w:t")
        node.set(W_NS + "space", "preserve")
        node.text = line
        run.append(node)
        paragraph.append(run)
        tc.append(paragraph)


def set_row_value(row, text):
    cells = uniq(row)
    if len(cells) > 1:
        set_cell_text(cells[1], text)
    elif cells:
        set_cell_text(cells[0], text)


def set_io(row, entradas, salidas):
    cells = uniq(row)
    if len(cells) >= 2:
        set_cell_text(cells[0], entradas)
        set_cell_text(cells[1], salidas)


def fill_cu_table(table, hu, nombre=None, desc=None, entradas=None, salidas=None, criterios=None, rol=None):
    if len(table.rows) < 8:
        return
    ent, sal = entradas_salidas(hu)
    set_row_value(table.rows[0], nombre or f"{hu['cu']} — {hu['titulo']}")
    set_row_value(table.rows[1], f"{hu['cu']} · {hu['rf']}")
    set_row_value(table.rows[2], desc or hu["desc_rf"])
    set_io(table.rows[4], entradas or ent, salidas or sal)
    set_row_value(table.rows[5], criterios or criterios_texto(hu))
    set_row_value(table.rows[6], rol or hu["actor"])
    set_row_value(table.rows[7], prio_word(hu["moscow_raw"]))


def first_row_text(table):
    return " ".join(cell.text for cell in uniq(table.rows[0]))


def ident_text(table):
    if len(table.rows) < 2:
        return ""
    return " ".join(cell.text for cell in uniq(table.rows[1]))


def fold(text):
    return (
        (text or "")
        .lower()
        .replace("á", "a")
        .replace("é", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ú", "u")
    )


JUNK = (
    "ayuda para el equipo",
    "sesiones: hay dos puertas",
    "escribe alex",
    "escribe sofia",
    "escribe paula",
    "escribe maicol",
    "revisa henry",
    "revisa alex",
    "texto del equipo",
    "copiar de la hoja",
    "completar desde",
    "ayuda: los requisitos no funcionales",
    "segunda redaccion",
    "apoyo que ya habian escrito",
    "en el kanban es",
    "igual que el kanban",
    "quien escribe",
    "paula y maicol no duplican",
    "no se borro lo que ya escribieron",
    "faltaba",
    "completar el resto",
)


def is_junk(text):
    compact = fold(text)
    return any(flag in compact for flag in JUNK)


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for extra in paragraph.runs[1:]:
            extra.text = ""
        return
    paragraph.add_run(text)


def build_word(hus):
    if not SRC_WORD.exists():
        raise SystemExit(f"No está la plantilla SENA: {SRC_WORD}")
    doc = Document(str(SRC_WORD))
    by_cu = {hu["cu"].replace(" ", "").upper(): hu for hu in hus}

    for row in doc.tables[1].rows:
        text = " ".join(cell.text for cell in uniq(row))
        if text.strip().startswith("Fecha:"):
            set_cell_text(uniq(row)[0], "Fecha: 23 de agosto de 2026")
        if "Módulo:" in text or "Modulo:" in text:
            cells = uniq(row)
            if len(cells) >= 2:
                set_cell_text(cells[0], "Módulo: SIGEVA (urna, mesa y red)")
                set_cell_text(cells[1], "Submódulo: Especificación Sprint 1")

    entregables = [
        ("Componente", "Descripción"),
        ("Requerimientos", "Inventario de requisitos funcionales y no funcionales del producto (02-Requerimientos.xlsx)."),
        ("Casos de uso", "Tablas en formato SENA, letra Arial y color verde 92d050. Una historia = un caso de uso = un requisito funcional."),
        ("Historias de usuario", "Un cuadro con una fila por historia: como / quiero / para, qué entra, qué no entra y criterios de aceptación."),
        ("Backlog", "Lista priorizada de las 22 historias de producto para un centro piloto."),
        ("Código fuente y plataforma", "Entrega de construcción. No aplica a este sprint."),
        ("Manuales y pruebas", "Entrega de construcción. No aplica a este sprint."),
        ("Transferencia de conocimiento", "Entrega de construcción. No aplica a este sprint."),
    ]
    table_ent = doc.tables[2]
    for index, (left, right) in enumerate(entregables):
        if index >= len(table_ent.rows):
            break
        cells = uniq(table_ent.rows[index])
        if len(cells) >= 2:
            set_cell_text(cells[0], left)
            set_cell_text(cells[1], right)

    drop = []
    for table in doc.tables:
        nombre = first_row_text(table)
        ident = ident_text(table)
        blob = nombre + " " + ident
        low = fold(blob)
        if "kanban" in low or "quien escribe" in low:
            drop.append(table._tbl)
            continue
        if "redaccion adicional" in low:
            drop.append(table._tbl)
            continue
        if "controlar periodo" in low or "controlar período" in fold(blob):
            continue
        if not nombre.strip().lower().startswith("nombre del requerimiento"):
            continue
        if any(flag in low for flag in ("seguridad de acceso", "proteccion de la", "disponibilidad", "facilidad de uso")):
            continue
        match = re.search(r"CU-(\d+)", blob)
        if not match:
            continue
        key = f"CU-{int(match.group(1)):02d}"
        hu = by_cu.get(key)
        if not hu:
            continue
        extra = {}
        if "generar acta" in low:
            extra = {
                "nombre": "CU-18 — Generar acta de resultados",
                "desc": "El sistema deberá permitir al Funcionario generar un documento PDF con la información principal y los resultados finales de una elección de su centro.",
                "entradas": "Elección del centro del Funcionario.",
                "salidas": "Documento con resultados por candidato, participación, ganador o empate.",
                "rol": "Funcionario",
            }
        elif key == "CU-18":
            extra = {"nombre": "CU-18 — Consultar resultados", "rol": "Funcionario"}
        if key == "CU-06":
            extra["rol"] = "Aprendiz (Votante)"
        if key == "CU-04":
            extra["desc"] = "El sistema deberá enviar un código de validación de 6 caracteres al correo del censo del Aprendiz. Enviar el código no registra el voto."
        if key == "CU-05":
            extra["desc"] = "El sistema deberá validar el código de 6 caracteres enviado al correo. Validar el código no registra el voto."
        fill_cu_table(table, hu, **extra)

    for tbl in drop:
        parent = tbl.getparent()
        if parent is not None:
            parent.remove(tbl)

    glossary = {
        "aprendiz: en el kanban": "Aprendiz: persona habilitada para votar (en este documento también se llama votante).",
        "funcionario: en el kanban": "Funcionario: usuario que administra los procesos de elección de un centro (gestor de elección).",
        "otp / codigo de validacion": "Código de validación: código de un solo uso enviado al correo del censo. Solicitarlo es CU-04. Validarlo es CU-05. Validar no es votar.",
    }
    body = doc.element.body
    for child in list(body):
        tag = child.tag.split("}")[-1]
        if tag != "p":
            continue
        raw = "".join(child.itertext())
        compact = " ".join(raw.split())
        if not compact:
            continue
        key = fold(compact)
        replaced = False
        for needle, nuevo in glossary.items():
            if key.startswith(needle):
                for paragraph in doc.paragraphs:
                    if paragraph._p is child:
                        set_paragraph_text(paragraph, nuevo)
                        replaced = True
                        break
        if replaced:
            continue
        if compact.upper().startswith("CU-") and is_junk(compact):
            limpio = re.split(r"FALTABA|FALTA|TEXTO DEL EQUIPO|Completar|Faltaba", compact, maxsplit=1)[0].strip(" .")
            for paragraph in doc.paragraphs:
                if paragraph._p is child:
                    set_paragraph_text(paragraph, limpio)
                    break
            continue
        if is_junk(compact):
            parent = child.getparent()
            if parent is not None:
                parent.remove(child)

    dest = OUT / "03-Casos-de-Uso.docx"
    doc.save(str(dest))
    print("OK", dest.name)


def write_readme():
    (OUT / "README.md").write_text(
        """# Sprint 1 — entregables SIGEVA

Semana del 24 al 28 de agosto de 2026. Se especifica el producto. No se construye código nuevo.

| Orden | Entregable | Archivo |
| --- | --- | --- |
| 1 | Sprint (qué entregar) | 01-Sprint.xlsx |
| 2 | Requerimientos | 02-Requerimientos.xlsx |
| 3 | Casos de uso (formato SENA) | 03-Casos-de-Uso.docx |
| 4 | Historias de usuario | 04-Historias-de-Usuario.xlsx |
| 5 | Backlog | 05-Backlog.xlsx |

Los casos de uso van en letra Arial y color verde 92d050, como la plantilla SENA.

Prioridad en español: Debe tener / Debería tener.

Una historia de producto = un requisito funcional = un caso de uso.
""",
        encoding="utf-8",
    )


def write_doc_readme():
    (HERE / "README.md").write_text(
        """# SIGEVA — documentación

Los entregables del Sprint 1 están en `entregables/sprint-1/`.

Ahí van requerimientos, casos de uso, historias de usuario y backlog.
""",
        encoding="utf-8",
    )


def borrar_viejos():
    for name in [
        "01-comprension-del-producto.md",
        "02-actores-y-personas.md",
        "03-journeys.md",
        "04-modulos.md",
        "05-epicas.md",
        "06-features.md",
        "08-casos-de-uso.md",
        "09-reglas-de-negocio.md",
        "10-requisitos-no-funcionales.md",
        "11-matriz-permisos.md",
        "12-entidades-y-datos.md",
        "13-backlog.md",
        "14-trazabilidad.md",
        "15-mvp.md",
        "16-sprints.md",
        "17-definition-of-ready-and-done.md",
        "18-riesgos.md",
        "19-auditoria-senior.md",
        "20-preguntas-pendientes.md",
    ]:
        path = HERE / name
        if path.exists():
            path.unlink()
            print("BORRADO", path.name)

    hu_dir = HERE / "07-historias-de-usuario"
    if hu_dir.exists():
        shutil.rmtree(hu_dir)
        print("BORRADO", hu_dir.name)

    for path in list(HERE.glob("SIGEVA-*.xlsx")) + list(HERE.glob("SIGEVA-*.docx")):
        path.unlink()
        print("BORRADO", path.name)

    for name in ("_excel_cuadro.py", "_excel_senior_generar.py", "_generar_office.py", "_gen_dos_excel.py"):
        path = HERE / name
        if path.exists():
            path.unlink()
            print("BORRADO", path.name)

    docs = REPO / "docs"
    if docs.exists():
        for path in docs.iterdir():
            if path.is_file() and path.suffix.lower() in {".xlsx", ".docx", ".md"}:
                path.unlink()
                print("BORRADO docs/", path.name)

    tmp = REPO / ".tmp-gen-cu-keep-format.py"
    if tmp.exists():
        tmp.unlink()
        print("BORRADO", tmp.name)

    cache = HERE / "__pycache__"
    if cache.exists():
        shutil.rmtree(cache)


def main():
    hus = all_hus()
    OUT.mkdir(parents=True, exist_ok=True)
    build_sprint().save(OUT / "01-Sprint.xlsx")
    print("OK 01-Sprint.xlsx")
    build_requerimientos(hus).save(OUT / "02-Requerimientos.xlsx")
    print("OK 02-Requerimientos.xlsx")
    build_word(hus)
    build_historias(hus).save(OUT / "04-Historias-de-Usuario.xlsx")
    print("OK 04-Historias-de-Usuario.xlsx")
    build_backlog(hus).save(OUT / "05-Backlog.xlsx")
    print("OK 05-Backlog.xlsx")
    write_readme()
    write_doc_readme()
    borrar_viejos()
    print("LISTO", OUT)


if __name__ == "__main__":
    main()
