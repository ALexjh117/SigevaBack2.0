# -*- coding: utf-8 -*-
"""Estilos institucionales para los Excel senior de SIGEVA."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

NAVY = "1B3A4B"
GREEN = "1F7A4D"
GOLD = "9A6B2F"
RED = "8B1E3F"
AMBER_BG = "FEF3C7"
AMBER_TX = "92400E"
OK_BG = "DCFCE7"
OK_TX = "166534"
NO_BG = "FEE2E2"
NO_TX = "991B1B"
GRAY = "F2F4F7"
WHITE = "FFFFFF"
MUTED = "667085"
LINE = "D0D5DD"
LIGHT_NAVY = "E8EEF2"
LIGHT_GREEN = "E8F5EE"
LIGHT_GOLD = "F8F1E3"

THIN = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)


def pf(h):
    return PatternFill("solid", fgColor=h)


def font(size=10, bold=False, color=NAVY, italic=False):
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)


def align(wrap=True, v="top", h="left"):
    return Alignment(wrap_text=wrap, vertical=v, horizontal=h)


def widths(ws, xs):
    for i, x in enumerate(xs, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def paint(cell, value=None, fill=None, fnt=None, al=None, border=THIN):
    if value is not None:
        cell.value = value
    if fill:
        cell.fill = pf(fill)
    cell.font = fnt or font()
    cell.alignment = al or align()
    cell.border = border
    return cell


def merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.5, header=0.25, footer=0.25)
    ws.sheet_view.showGridLines = False
    ws.oddHeader.left.text = "SIGEVA  ·  SENA"
    ws.oddHeader.right.text = "&A"
    ws.oddFooter.left.text = "Confidencial — uso interno del equipo de producto"
    ws.oddFooter.right.text = "Pág. &P de &N"


def header_bar(ws, row, n, color=NAVY, height=28):
    for c in range(1, n + 1):
        cell = ws.cell(row, c)
        cell.fill = pf(color)
        cell.font = font(10, True, WHITE)
        cell.alignment = align(True, "center", "center")
        cell.border = THIN
    ws.row_dimensions[row].height = height


def freeze_filter(ws, header_row, n_cols, last_row):
    ws.freeze_panes = f"A{header_row + 1}"
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(n_cols)}{last_row}"


def moscow_fill(v):
    v = (v or "").upper()
    if v.startswith("MUST"):
        return GREEN, WHITE
    if v.startswith("SHOULD"):
        return GOLD, WHITE
    if v.startswith("COULD"):
        return MUTED, WHITE
    if "WON'T" in v or "WONT" in v:
        return RED, WHITE
    return GRAY, NAVY


def dor_fill(v):
    v = (v or "").upper()
    if "NO READY" in v or "BLOQUE" in v:
        return NO_BG, NO_TX
    if "PARCIAL" in v or "CONDICION" in v:
        return AMBER_BG, AMBER_TX
    if "READY" in v or v.startswith("SÍ") or v.startswith("SI "):
        return OK_BG, OK_TX
    return GRAY, NAVY


def tipo_ca_fill(v):
    v = (v or "").upper()
    if "FELIZ" in v or "PRINCIPAL" in v:
        return OK_BG, OK_TX
    if "PENDIENTE" in v or "CONDICION" in v:
        return AMBER_BG, AMBER_TX
    if "ERROR" in v or "EXCEPC" in v or "SEGUR" in v or "PERMISO" in v:
        return NO_BG, NO_TX
    return LIGHT_NAVY, NAVY
